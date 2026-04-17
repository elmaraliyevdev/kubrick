import json
import logging
import re
from typing import Dict, List, Optional, Any, Union
from dataclasses import dataclass
from datetime import datetime
import openai
from openai import OpenAI
import asyncio
from enum import Enum

logger = logging.getLogger(__name__)


class ParsingStrategy(Enum):
    """Different parsing strategies based on resume complexity"""
    SIMPLE = "simple"
    DETAILED = "detailed" 
    MULTILINGUAL = "multilingual"
    TECHNICAL = "technical"


@dataclass
class ParsingConfig:
    """Configuration for LLM-based parsing"""
    model: str = "gpt-3.5-turbo"
    temperature: float = 0.1
    max_tokens: int = 2000
    timeout: int = 30
    retry_attempts: int = 3
    fallback_to_regex: bool = False


class LLMResumeParser:
    """
    Advanced LLM-based resume parser that can handle diverse formats, 
    languages, and industries with intelligent contextual understanding.
    """
    
    def __init__(self, config: ParsingConfig = None, api_key: str = None):
        self.config = config or ParsingConfig()
        
        # Initialize OpenAI client with minimal configuration
        self.client = None
        try:
            import os
            # Get API key from parameter or environment
            openai_api_key = api_key or os.getenv('OPENAI_API_KEY')
            
            if openai_api_key and openai_api_key != 'your_openai_api_key_here':
                # Store API key for direct requests to avoid client issues
                self.api_key = openai_api_key
                self.client = None  # We'll use direct API calls
                logger.info("OpenAI API key configured for direct requests")
            else:
                raise Exception("No valid OpenAI API key found. LLM parser requires OPENAI_API_KEY environment variable.")
                
        except Exception as e:
            logger.error(f"OpenAI client initialization failed: {e}")
            raise Exception(f"LLM parser requires valid OpenAI API key. Error: {e}")
                
        # No fallback parser - LLM only
        self.fallback_parser = None

    def _detect_language(self, text: str) -> str:
        """Detect the primary language of the resume"""
        # Simple language detection based on common words
        languages = {
            'english': ['experience', 'education', 'skills', 'work', 'university', 'degree'],
            'spanish': ['experiencia', 'educación', 'habilidades', 'trabajo', 'universidad'],
            'french': ['expérience', 'éducation', 'compétences', 'travail', 'université'],
            'german': ['erfahrung', 'bildung', 'fähigkeiten', 'arbeit', 'universität'],
            'russian': ['опыт', 'образование', 'навыки', 'работа', 'университет'],
            'chinese': ['经验', '教育', '技能', '工作', '大学'],
        }
        
        text_lower = text.lower()
        scores = {}
        
        for lang, keywords in languages.items():
            score = sum(1 for keyword in keywords if keyword in text_lower)
            scores[lang] = score
            
        detected = max(scores, key=scores.get) if scores else 'english'
        return detected

    def _determine_parsing_strategy(self, text: str) -> ParsingStrategy:
        """Determine the best parsing strategy based on resume content"""
        text_lower = text.lower()
        
        # Technical resume indicators
        tech_keywords = ['software', 'developer', 'engineer', 'programming', 'python', 'javascript', 
                        'react', 'node.js', 'api', 'database', 'cloud', 'aws', 'docker']
        tech_score = sum(1 for keyword in tech_keywords if keyword in text_lower)
        
        # Language indicators
        is_multilingual = len([lang for lang in ['english', 'spanish', 'french', 'german', 'russian', 'chinese'] 
                              if any(word in text_lower for word in [lang, lang[:3]])]) > 1
        
        # Complexity indicators
        is_complex = len(text) > 3000 or text.count('\n') > 50
        
        if is_multilingual:
            return ParsingStrategy.MULTILINGUAL
        elif tech_score > 5:
            return ParsingStrategy.TECHNICAL
        elif is_complex:
            return ParsingStrategy.DETAILED
        else:
            return ParsingStrategy.SIMPLE

    def _build_parsing_prompt(self, text: str, strategy: ParsingStrategy, language: str) -> str:
        """Build contextual prompt based on strategy and language"""
        
        base_structure = {
            "personal_info": {
                "full_name": "string",
                "email": "string", 
                "phone": "string",
                "location": "string (city, country)",
                "linkedin": "string (URL if available)",
                "website": "string (URL if available)"
            },
            "work_experience": [
                {
                    "company_name": "string",
                    "job_title": "string", 
                    "start_date": "string (MM/YYYY format)",
                    "end_date": "string (MM/YYYY format or 'Present')",
                    "location": "string (if mentioned)",
                    "description": "string (key responsibilities and achievements)"
                }
            ],
            "education": [
                {
                    "school": "string (university/school name)",
                    "degree": "string (degree type and field)",
                    "start_date": "string (MM/YYYY if available)",
                    "end_date": "string (MM/YYYY format)",
                    "location": "string (if mentioned)",
                    "gpa": "string (if mentioned)"
                }
            ],
            "skills": {
                "technical_skills": ["list of technical skills"],
                "soft_skills": ["list of soft skills"],
                "languages": ["list of spoken languages with proficiency if mentioned"],
                "certifications": ["list of certifications"]
            },
            "projects": [
                {
                    "name": "string",
                    "description": "string", 
                    "technologies": ["list of technologies used"],
                    "date": "string (if available)"
                }
            ]
        }
        
        language_instructions = {
            'english': 'The resume is in English.',
            'spanish': 'The resume is in Spanish. Please extract information and provide output in English.',
            'french': 'The resume is in French. Please extract information and provide output in English.',
            'german': 'The resume is in German. Please extract information and provide output in English.',
            'russian': 'The resume is in Russian. Please extract information and provide output in English.',
            'chinese': 'The resume is in Chinese. Please extract information and provide output in English.',
        }
        
        strategy_instructions = {
            ParsingStrategy.SIMPLE: "Focus on extracting basic information clearly and accurately.",
            ParsingStrategy.DETAILED: "Extract comprehensive details including all mentioned achievements, metrics, and specific technologies.",
            ParsingStrategy.MULTILINGUAL: "Pay special attention to language sections and international experience. Handle mixed languages appropriately.",
            ParsingStrategy.TECHNICAL: "Focus on technical skills, programming languages, frameworks, tools, and technical project details. Extract version numbers and specific technologies when mentioned."
        }
        
        prompt = f"""You are an expert resume parser. {language_instructions.get(language, '')} {strategy_instructions[strategy]}

Please extract information from the following resume and return it as valid JSON matching this exact structure:

{json.dumps(base_structure, indent=2)}

IMPORTANT GUIDELINES:
1. Return ONLY valid JSON, no additional text or explanations
2. If information is not found, use null instead of empty strings
3. For dates, always use MM/YYYY format (e.g., "03/2021") 
4. For current positions, use "Present" as end_date
5. Extract ALL work experiences and education entries, don't skip any
6. Be flexible with job titles and company names - extract exactly as written
7. For skills, categorize appropriately (technical vs soft skills)
8. If a field contains multiple values, use arrays as shown in the structure
9. For descriptions, capture key achievements and responsibilities concisely
10. Handle various resume formats (chronological, functional, combination)

Resume text:
{text}
"""
        
        return prompt

    def _call_llm(self, prompt: str) -> Optional[Dict[str, Any]]:
        """Make API call to LLM using direct HTTP requests"""
        if not hasattr(self, 'api_key') or not self.api_key:
            return None
            
        import random
        import requests
        import time
        
        headers = {
            'Authorization': f'Bearer {self.api_key}',
            'Content-Type': 'application/json'
        }
        
        data = {
            'model': self.config.model,
            'messages': [
                {"role": "system", "content": "You are a professional resume parsing assistant. Always return valid JSON."},
                {"role": "user", "content": prompt}
            ],
            'temperature': self.config.temperature,
            'max_tokens': self.config.max_tokens
        }
            
        for attempt in range(self.config.retry_attempts):
            try:
                response = requests.post(
                    'https://api.openai.com/v1/chat/completions',
                    headers=headers,
                    json=data,
                    timeout=(10, self.config.timeout),
                )
                
                if response.status_code == 200:
                    response_data = response.json()
                    content = response_data['choices'][0]['message']['content'].strip()
                    
                    # Clean up response (remove markdown code blocks if present)
                    if content.startswith('```json'):
                        content = content[7:]
                    if content.startswith('```'):
                        content = content[3:]
                    if content.endswith('```'):
                        content = content[:-3]
                    
                    # Parse JSON
                    result = json.loads(content.strip())
                    return result
                else:
                    logger.warning(f"OpenAI API error {response.status_code}: {response.text}")
                    retryable_status = response.status_code in {408, 409, 429, 500, 502, 503, 504}
                    if not retryable_status or attempt == self.config.retry_attempts - 1:
                        logger.error("Failed to call OpenAI API after all retries")
                        return None
                
            except json.JSONDecodeError as e:
                logger.warning(f"Invalid JSON response on attempt {attempt + 1}: {e}")
                if attempt == self.config.retry_attempts - 1:
                    logger.error("Failed to get valid JSON after all retries")
                    return None
                    
            except Exception as e:
                logger.warning(f"LLM API call failed on attempt {attempt + 1}: {e}")
                if attempt == self.config.retry_attempts - 1:
                    logger.error("Failed to call LLM after all retries")
                    return None
                    
            # Jitter prevents synchronized retry spikes across concurrent requests.
            sleep_seconds = (2 ** attempt) + random.uniform(0, 0.25)
            time.sleep(sleep_seconds)
                
        return None

    def _clean_and_validate_result(self, result: Dict[str, Any]) -> Dict[str, Any]:
        """Clean and validate the LLM parsing result"""
        if not result:
            return {}
            
        # Standardize date formats
        def clean_date(date_str):
            if not date_str or date_str.lower() in ['null', 'none', '']:
                return None
            if date_str.lower() == 'present':
                return 'Present'
            
            # Try to standardize to MM/YYYY format
            date_patterns = [
                r'(\d{1,2})/(\d{4})',  # MM/YYYY or M/YYYY
                r'(\d{4})-(\d{1,2})',  # YYYY-MM
                r'(\d{1,2})-(\d{4})',  # MM-YYYY
                r'(\d{4})',            # Just year
            ]
            
            for pattern in date_patterns:
                match = re.search(pattern, str(date_str))
                if match:
                    if len(match.groups()) == 2:
                        month, year = match.groups()
                        if len(year) == 4:  # Year is second
                            return f"{int(month):02d}/{year}"
                        else:  # Year is first
                            return f"{int(match.groups()[1]):02d}/{match.groups()[0]}"
                    else:  # Just year
                        return f"01/{match.group(1)}"
            
            return date_str  # Return as-is if no pattern matches
        
        # Clean work experience
        if 'work_experience' in result and isinstance(result['work_experience'], list):
            for exp in result['work_experience']:
                if isinstance(exp, dict):
                    exp['start_date'] = clean_date(exp.get('start_date'))
                    exp['end_date'] = clean_date(exp.get('end_date'))
                    
                    # Ensure required fields
                    exp.setdefault('company_name', None)
                    exp.setdefault('job_title', None)
                    exp.setdefault('description', None)
        
        # Clean education
        if 'education' in result and isinstance(result['education'], list):
            for edu in result['education']:
                if isinstance(edu, dict):
                    edu['start_date'] = clean_date(edu.get('start_date'))
                    edu['end_date'] = clean_date(edu.get('end_date'))
                    
                    # For education, extract year from end_date if available
                    if edu.get('end_date') and edu['end_date'] != 'Present':
                        year_match = re.search(r'(\d{4})', str(edu['end_date']))
                        if year_match:
                            edu['year'] = int(year_match.group(1))
        
        # Clean skills
        if 'skills' in result and isinstance(result['skills'], dict):
            for skill_type in ['technical_skills', 'soft_skills', 'languages', 'certifications']:
                if skill_type in result['skills']:
                    skills = result['skills'][skill_type]
                    if isinstance(skills, list):
                        # Remove empty/null skills
                        result['skills'][skill_type] = [s for s in skills if s and str(s).strip()]
                    elif isinstance(skills, str):
                        # Convert single string to list
                        result['skills'][skill_type] = [s.strip() for s in skills.split(',') if s.strip()]
        
        return result

    async def parse_resume_with_llm(self, text: str) -> Dict[str, Any]:
        """Main LLM-based parsing method"""
        try:
            # Detect language and strategy
            language = self._detect_language(text)
            strategy = self._determine_parsing_strategy(text)
            
            logger.info(f"Parsing resume with strategy: {strategy.value}, language: {language}")
            
            # Build prompt
            prompt = self._build_parsing_prompt(text, strategy, language)
            
            # Call LLM in thread to avoid blocking
            import concurrent.futures
            
            loop = asyncio.get_event_loop()
            with concurrent.futures.ThreadPoolExecutor() as executor:
                result = await loop.run_in_executor(executor, self._call_llm, prompt)
            
            if result:
                # Clean and validate
                cleaned_result = self._clean_and_validate_result(result)
                
                # Add metadata
                cleaned_result['parsing_metadata'] = {
                    'method': 'llm',
                    'strategy': strategy.value,
                    'language': language,
                    'model': self.config.model,
                    'parsed_at': datetime.utcnow().isoformat()
                }
                
                return cleaned_result
            else:
                raise Exception("LLM parsing failed")
                
        except Exception as e:
            logger.error(f"LLM parsing error: {e}")
            
            # No fallback - raise the error
            raise Exception(f"LLM parsing failed: {str(e)}")

    async def parse_resume(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """Main entry point - extract text and parse with LLM"""
        try:
            # Extract text from file
            text = await self._extract_text_from_file(file_content, filename)
            
            if not text or len(text.strip()) < 50:
                return {
                    'error': 'Insufficient text content',
                    'parsing_metadata': {'method': 'failed', 'error': 'No extractable text'}
                }
            
            # Parse with LLM
            return await self.parse_resume_with_llm(text)
            
        except Exception as e:
            logger.error(f"Resume parsing failed: {e}")
            return {
                'error': 'File processing failed', 
                'details': str(e),
                'parsing_metadata': {'method': 'failed', 'error': str(e)}
            }

    async def _extract_text_from_file(self, file_content: bytes, filename: str) -> str:
        """Extract text from various file formats"""
        filename_lower = filename.lower()
        
        try:
            if filename_lower.endswith('.pdf'):
                # PDF extraction
                import io
                import PyPDF2
                
                pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
                text = ""
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
                return text
                
            elif filename_lower.endswith(('.docx', '.doc')):
                # Word document extraction
                import io
                import docx
                
                doc = docx.Document(io.BytesIO(file_content))
                text = ""
                for paragraph in doc.paragraphs:
                    text += paragraph.text + "\n"
                return text
                
            elif filename_lower.endswith('.txt'):
                # Plain text
                return file_content.decode('utf-8', errors='ignore')
                
            else:
                # Try to decode as text
                return file_content.decode('utf-8', errors='ignore')
                
        except Exception as e:
            logger.error(f"Text extraction failed: {e}")
            raise Exception(f"Could not extract text from {filename}: {str(e)}")


class RegexResumeParser:
    """
    Regex-based resume parser used as a fallback when the LLM parser
    is unavailable (e.g. no OpenAI API key configured).
    Returns data in the same structure as LLMResumeParser.
    """

    # Common email pattern
    EMAIL_RE = re.compile(r'[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}')
    # Phone patterns (international)
    PHONE_RE = re.compile(r'[\+]?[\d\s\-\(\)\.]{7,20}')
    # LinkedIn URL
    LINKEDIN_RE = re.compile(r'(?:https?://)?(?:www\.)?linkedin\.com/in/[\w\-]+/?', re.IGNORECASE)
    # Generic URL
    URL_RE = re.compile(r'https?://[^\s,]+')
    # Section headers
    SECTION_RE = re.compile(
        r'^[\s]*(?:EXPERIENCE|WORK\s*EXPERIENCE|EMPLOYMENT|PROFESSIONAL\s*EXPERIENCE'
        r'|EDUCATION|ACADEMIC|SKILLS|TECHNICAL\s*SKILLS|CORE\s*COMPETENCIES'
        r'|PROJECTS|CERTIFICATIONS?|LANGUAGES|SUMMARY|OBJECTIVE|PROFILE'
        r'|AWARDS|HONORS|PUBLICATIONS|VOLUNTEER)'
        r'[\s:]*$',
        re.IGNORECASE | re.MULTILINE
    )
    # Date patterns like "Jan 2020", "01/2020", "2020", "Present"
    DATE_RE = re.compile(
        r'(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?'
        r'|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)'
        r'[\s.,\-]*\d{4}'
        r'|\d{1,2}[/\-]\d{4}'
        r'|\d{4}'
        r'|[Pp]resent|[Cc]urrent',
        re.IGNORECASE
    )
    # Date range like "Jan 2020 - Present" or "2019 – 2021"
    DATE_RANGE_RE = re.compile(
        r'('
        r'(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?'
        r'|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)'
        r'[\s.,]*\d{4}'
        r'|\d{1,2}[/\-]\d{4}'
        r'|\d{4}'
        r')'
        r'\s*[\-–—~to]+\s*'
        r'('
        r'(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?'
        r'|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)'
        r'[\s.,]*\d{4}'
        r'|\d{1,2}[/\-]\d{4}'
        r'|\d{4}'
        r'|[Pp]resent|[Cc]urrent'
        r')',
        re.IGNORECASE
    )

    # Common technical skills for matching
    TECH_SKILLS = {
        'python', 'java', 'javascript', 'typescript', 'c++', 'c#', 'ruby', 'go', 'golang',
        'rust', 'swift', 'kotlin', 'php', 'scala', 'r', 'matlab', 'perl', 'sql', 'html',
        'css', 'react', 'angular', 'vue', 'vue.js', 'node.js', 'nodejs', 'express',
        'django', 'flask', 'fastapi', 'spring', 'spring boot', '.net', 'rails',
        'docker', 'kubernetes', 'k8s', 'aws', 'azure', 'gcp', 'terraform', 'ansible',
        'jenkins', 'git', 'github', 'gitlab', 'ci/cd', 'linux', 'unix', 'bash',
        'postgresql', 'mysql', 'mongodb', 'redis', 'elasticsearch', 'dynamodb',
        'kafka', 'rabbitmq', 'graphql', 'rest', 'restful', 'api', 'microservices',
        'machine learning', 'deep learning', 'tensorflow', 'pytorch', 'scikit-learn',
        'pandas', 'numpy', 'spark', 'hadoop', 'tableau', 'power bi',
        'figma', 'sketch', 'adobe', 'photoshop', 'illustrator',
        'agile', 'scrum', 'jira', 'confluence', 'devops',
    }

    async def _extract_text_from_file(self, file_content: bytes, filename: str) -> str:
        """Extract text from various file formats (same as LLMResumeParser)."""
        filename_lower = filename.lower()
        try:
            if filename_lower.endswith('.pdf'):
                import io, PyPDF2
                pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
                return "\n".join(page.extract_text() for page in pdf_reader.pages)
            elif filename_lower.endswith(('.docx', '.doc')):
                import io, docx
                doc = docx.Document(io.BytesIO(file_content))
                return "\n".join(p.text for p in doc.paragraphs)
            else:
                return file_content.decode('utf-8', errors='ignore')
        except Exception as e:
            logger.error(f"Text extraction failed: {e}")
            raise Exception(f"Could not extract text from {filename}: {str(e)}")

    def _split_sections(self, text: str) -> Dict[str, str]:
        """Split resume text into sections by common headings."""
        sections: Dict[str, str] = {}
        header_positions = []
        for m in self.SECTION_RE.finditer(text):
            header_positions.append((m.start(), m.end(), m.group().strip().upper()))

        if not header_positions:
            return {"FULL": text}

        # Text before first header is the "HEADER" section (name, contacts)
        if header_positions[0][0] > 0:
            sections["HEADER"] = text[:header_positions[0][0]]

        for i, (start, end, name) in enumerate(header_positions):
            next_start = header_positions[i + 1][0] if i + 1 < len(header_positions) else len(text)
            # Normalize similar headers
            if any(k in name for k in ("EXPERIENCE", "EMPLOYMENT")):
                name = "EXPERIENCE"
            elif any(k in name for k in ("EDUCATION", "ACADEMIC")):
                name = "EDUCATION"
            elif any(k in name for k in ("SKILL", "COMPETENC")):
                name = "SKILLS"
            elif "PROJECT" in name:
                name = "PROJECTS"
            sections[name] = text[end:next_start]

        return sections

    def _extract_personal_info(self, text: str, header_text: str) -> Dict[str, Optional[str]]:
        """Extract personal info from the resume header area."""
        search_area = header_text or text[:500]

        email_match = self.EMAIL_RE.search(search_area)
        email = email_match.group() if email_match else None

        linkedin_match = self.LINKEDIN_RE.search(search_area)
        linkedin = linkedin_match.group() if linkedin_match else None

        # Find website URL that is not linkedin
        website = None
        for url_match in self.URL_RE.finditer(search_area):
            url = url_match.group()
            if 'linkedin.com' not in url:
                website = url
                break

        # Phone — grab first match that looks like a phone number
        phone = None
        for phone_match in self.PHONE_RE.finditer(search_area):
            candidate_phone = phone_match.group().strip()
            # Must have at least 7 digits
            if sum(c.isdigit() for c in candidate_phone) >= 7:
                phone = candidate_phone
                break

        # Name: typically the first non-empty line (before email/phone)
        full_name = None
        lines = [l.strip() for l in search_area.split('\n') if l.strip()]
        for line in lines[:5]:
            # Skip lines that look like email, phone, URL, or section headers
            if self.EMAIL_RE.search(line) or self.LINKEDIN_RE.search(line):
                continue
            if self.URL_RE.search(line):
                continue
            if self.SECTION_RE.match(line):
                continue
            if sum(c.isdigit() for c in line) > len(line) * 0.4:
                continue
            # Likely a name if it's short and mostly alphabetic
            if len(line) < 60 and sum(c.isalpha() or c.isspace() for c in line) > len(line) * 0.7:
                full_name = line
                break

        # Location: look for city/state/country patterns
        location = None
        location_re = re.compile(
            r'(?:^|\||\•|\,)\s*([A-Z][a-zA-Z\s]+,\s*[A-Z]{2,}(?:\s+\d{5})?)',
        )
        loc_match = location_re.search(search_area)
        if loc_match:
            location = loc_match.group(1).strip()

        return {
            "full_name": full_name,
            "email": email,
            "phone": phone,
            "location": location,
            "linkedin": linkedin,
            "website": website,
        }

    def _extract_work_experience(self, section_text: str) -> List[Dict[str, Optional[str]]]:
        """Extract work experience entries from the experience section."""
        if not section_text:
            return []

        entries: List[Dict[str, Optional[str]]] = []
        lines = section_text.strip().split('\n')

        current_entry: Optional[Dict[str, Optional[str]]] = None
        description_lines: List[str] = []

        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue

            date_range = self.DATE_RANGE_RE.search(stripped)
            # A new entry typically has a date range and is not a bullet point
            is_bullet = stripped.startswith(('•', '-', '–', '*', '◦', '▪'))

            if date_range and not is_bullet:
                # Save previous entry
                if current_entry:
                    current_entry['description'] = ' '.join(description_lines) if description_lines else None
                    entries.append(current_entry)
                    description_lines = []

                # Parse the line for company/title
                text_without_dates = stripped[:date_range.start()] + stripped[date_range.end():]
                parts = re.split(r'\s*[\|–—•·]\s*', text_without_dates.strip(' |-–—'))
                parts = [p.strip() for p in parts if p.strip()]

                job_title = parts[0] if len(parts) >= 1 else None
                company_name = parts[1] if len(parts) >= 2 else None
                location = parts[2] if len(parts) >= 3 else None

                current_entry = {
                    'job_title': job_title,
                    'company_name': company_name,
                    'location': location,
                    'start_date': date_range.group(1),
                    'end_date': date_range.group(2),
                    'description': None,
                }
            elif is_bullet and current_entry:
                description_lines.append(stripped.lstrip('•-–*◦▪ '))
            elif current_entry:
                # Could be continuation or company line
                if not current_entry.get('company_name'):
                    current_entry['company_name'] = stripped
                else:
                    description_lines.append(stripped)

        # Don't forget last entry
        if current_entry:
            current_entry['description'] = ' '.join(description_lines) if description_lines else None
            entries.append(current_entry)

        return entries

    def _extract_education(self, section_text: str) -> List[Dict[str, Optional[str]]]:
        """Extract education entries."""
        if not section_text:
            return []

        entries: List[Dict[str, Optional[str]]] = []
        lines = section_text.strip().split('\n')

        current_entry: Optional[Dict[str, Optional[str]]] = None

        for line in lines:
            stripped = line.strip()
            if not stripped or stripped.startswith(('•', '-', '*')):
                continue

            date_range = self.DATE_RANGE_RE.search(stripped)
            single_date = self.DATE_RE.search(stripped) if not date_range else None

            if date_range or single_date:
                if current_entry:
                    entries.append(current_entry)

                if date_range:
                    text_part = stripped[:date_range.start()] + stripped[date_range.end():]
                    start_date = date_range.group(1)
                    end_date = date_range.group(2)
                else:
                    text_part = stripped[:single_date.start()] + stripped[single_date.end():]
                    start_date = None
                    end_date = single_date.group()

                parts = re.split(r'\s*[\|–—•·,]\s*', text_part.strip(' |-–—,'))
                parts = [p.strip() for p in parts if p.strip()]

                school = None
                degree = None
                for p in parts:
                    if any(kw in p.lower() for kw in ('university', 'college', 'institute', 'school', 'academy')):
                        school = p
                    elif any(kw in p.lower() for kw in ('bachelor', 'master', 'phd', 'doctor', 'associate', 'degree', 'b.s', 'm.s', 'b.a', 'm.a', 'mba', 'diploma')):
                        degree = p
                if not school and parts:
                    school = parts[0]
                if not degree and len(parts) > 1:
                    degree = parts[1]

                current_entry = {
                    'school': school,
                    'degree': degree,
                    'start_date': start_date,
                    'end_date': end_date,
                    'location': None,
                    'gpa': None,
                }
            elif current_entry is None and stripped:
                # Line without date — might be the first education line
                current_entry = {
                    'school': stripped,
                    'degree': None,
                    'start_date': None,
                    'end_date': None,
                    'location': None,
                    'gpa': None,
                }

        if current_entry:
            entries.append(current_entry)

        return entries

    def _extract_skills(self, text: str, skills_section: str) -> Dict[str, List[str]]:
        """Extract skills from skills section and full text."""
        found_skills: List[str] = []
        search_text = (skills_section + '\n' + text).lower()

        for skill in self.TECH_SKILLS:
            # Word-boundary match
            pattern = r'(?:^|[\s,;|•\-/\(])' + re.escape(skill) + r'(?:[\s,;|•\-/\)]|$)'
            if re.search(pattern, search_text):
                found_skills.append(skill.title() if len(skill) > 3 else skill.upper())

        # Also grab comma/pipe-separated items from skills section
        if skills_section:
            for line in skills_section.split('\n'):
                stripped = line.strip().lstrip('•-–*◦▪ ')
                if not stripped:
                    continue
                items = re.split(r'[,;|•·]+', stripped)
                for item in items:
                    item = item.strip()
                    if 2 <= len(item) <= 40 and item.lower() not in {s.lower() for s in found_skills}:
                        found_skills.append(item)

        return {
            "technical_skills": found_skills,
            "soft_skills": [],
            "languages": [],
            "certifications": [],
        }

    async def parse_resume(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """Main entry point — extract text and parse with regex."""
        try:
            text = await self._extract_text_from_file(file_content, filename)

            if not text or len(text.strip()) < 50:
                return {
                    'personal_info': {},
                    'work_experience': [],
                    'education': [],
                    'skills': {},
                    'projects': [],
                    'raw_text': text,
                    'parsing_metadata': {'method': 'regex', 'error': 'Insufficient text content'},
                }

            sections = self._split_sections(text)
            header_text = sections.get("HEADER", "")

            personal_info = self._extract_personal_info(text, header_text)
            work_experience = self._extract_work_experience(sections.get("EXPERIENCE", ""))
            education = self._extract_education(sections.get("EDUCATION", ""))
            skills = self._extract_skills(text, sections.get("SKILLS", ""))

            return {
                'personal_info': personal_info,
                'work_experience': work_experience,
                'education': education,
                'skills': skills,
                'projects': [],
                'raw_text': text,
                'parsing_metadata': {
                    'method': 'regex',
                    'strategy': 'fallback',
                    'language': 'english',
                    'parsed_at': datetime.utcnow().isoformat(),
                },
            }

        except Exception as e:
            logger.error(f"Regex resume parsing failed: {e}")
            return {
                'personal_info': {},
                'work_experience': [],
                'education': [],
                'skills': {},
                'projects': [],
                'error': str(e),
                'parsing_metadata': {'method': 'failed', 'error': str(e)},
            }


# Factory function for easy integration
def create_llm_parser(api_key: str = None, **config_kwargs) -> LLMResumeParser:
    """Factory function to create LLM parser with custom config"""
    config = ParsingConfig(**config_kwargs)
    return LLMResumeParser(config=config, api_key=api_key)


# Example usage and testing
if __name__ == "__main__":
    async def test_parser():
        # Test with various resume formats
        test_resumes = [
            # English technical resume
            """
            John Doe
            john.doe@email.com | +1-555-0123
            San Francisco, CA
            
            EXPERIENCE
            Senior Software Engineer | Google | 2021 - Present
            • Developed scalable microservices using Python and Go
            • Led team of 5 engineers, improved deployment frequency by 300%
            
            Software Developer | Startup Inc | 2019 - 2021
            • Built React applications serving 1M+ users
            • Implemented CI/CD pipelines with Docker and Kubernetes
            
            EDUCATION
            Master of Science in Computer Science | Stanford University | 2017 - 2019
            Bachelor of Engineering | UC Berkeley | 2013 - 2017
            
            SKILLS
            Python, JavaScript, React, Docker, AWS, Machine Learning
            """,
            
            # Spanish resume
            """
            María García
            maria.garcia@correo.es | +34-600-123-456
            Madrid, España
            
            EXPERIENCIA PROFESIONAL
            Ingeniera de Software Senior | Telefónica | 2020 - Presente
            • Desarrollo de aplicaciones web con React y Node.js
            • Liderazgo de equipo de 3 desarrolladores
            
            Desarrolladora Junior | Banco Santander | 2018 - 2020
            • Mantenimiento de sistemas bancarios legacy
            
            EDUCACIÓN
            Ingeniería Informática | Universidad Politécnica de Madrid | 2014 - 2018
            
            HABILIDADES
            JavaScript, Python, React, SQL, Git
            """
        ]
        
        parser = create_llm_parser()
        
        for i, resume_text in enumerate(test_resumes, 1):
            print(f"\n=== Testing Resume {i} ===")
            result = await parser.parse_resume_with_llm(resume_text)
            print(json.dumps(result, indent=2))
    
    # Run tests
    asyncio.run(test_parser())